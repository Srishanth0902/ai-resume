import base64
import io
import json
import zipfile

import pymupdf as fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.detect import detect_pii
from app.extract import extract
from app.main import app
from app.redact import RedactItem, redact_document, redact_text

from .sample_resumes import KEEP, PII, make_docx, make_pdf

SAMPLES = {"pdf": make_pdf(), "docx": make_docx()}


def default_items(text):
    return [RedactItem(f.text, f.category) for f in detect_pii(text) if f.default_selected]


# --------------------------------------------------------------------------- detection


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_detects_personal_identifiers(kind):
    text = extract(SAMPLES[kind], kind).text
    selected = " || ".join(f.text.lower() for f in detect_pii(text) if f.default_selected)
    for value in PII:
        assert value.lower() in selected, f"missed {value!r}"


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_does_not_flag_skills(kind):
    text = extract(SAMPLES[kind], kind).text
    flagged = {f.text.lower() for f in detect_pii(text) if f.default_selected}
    for skill in ["python", "sql", "react", "fastapi", "docker", "postgresql", "chart.js", "gender"]:
        assert skill not in flagged


@pytest.mark.parametrize(
    "text, expected, category",
    [
        ("Contact: 98765 43210", "98765 43210", "phone"),
        ("Tel (040) 2345-6789", "(040) 2345-6789", "phone"),
        ("Aadhaar: 1234 5678 9012", "1234 5678 9012", "id_number"),
        ("PAN ABCDE1234F", "ABCDE1234F", "id_number"),
        ("DOB - 3rd March 2001", "3rd March 2001", "dob_age"),
        ("Age: 23", "23", "dob_age"),
        ("She led a team of five.", "She", "gender"),
        ("Mrs. Priya Sharma", "Mrs.", "gender"),
        ("Portfolio: https://priya.dev/work", "https://priya.dev/work", "url"),
    ],
)
def test_pattern_rules(text, expected, category):
    found = {(f.text, f.category) for f in detect_pii(text)}
    assert (expected, category) in found


def test_year_range_is_not_a_phone_number():
    assert not [f for f in detect_pii("B.Tech 2020 - 2024, 2019-2023") if f.category == "phone"]


# --------------------------------------------------------------------------- redaction


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_redacted_document_removes_pii_and_keeps_qualifications(kind):
    data = SAMPLES[kind]
    items = default_items(extract(data, kind).text)
    output, stats = redact_document(data, kind, items, "removed", remove_photos=True)
    out_text = extract(output, kind).text

    for value in PII:
        assert value.lower() not in out_text.lower(), f"{value!r} still present"
    for value in KEEP:
        assert value.lower() in out_text.lower(), f"{value!r} was removed"
    assert "[REMOVED]" in out_text
    assert stats.photos_removed == 1
    assert stats.links_removed >= 2


def test_pdf_output_has_no_hidden_leaks():
    data = SAMPLES["pdf"]
    output, _ = redact_document(data, "pdf", default_items(extract(data, "pdf").text), "removed", True)
    doc = fitz.open(stream=output, filetype="pdf")
    page = doc[0]
    assert not page.get_images(), "photo still embedded"
    assert all("asha" not in (link.get("uri") or "") for link in page.get_links())
    assert "asha" not in json.dumps(doc.metadata).lower()
    assert b"asha" not in output.lower()  # not hidden under a box: really gone from the file


def test_docx_output_has_no_hidden_leaks():
    data = SAMPLES["docx"]
    output, _ = redact_document(data, "docx", default_items(extract(data, "docx").text), "category", True)
    with zipfile.ZipFile(io.BytesIO(output)) as z:
        names = z.namelist()
        assert not [n for n in names if n.startswith("word/media/")], "photo still inside the package"
        blob = b"".join(z.read(n) for n in names).lower()
    assert b"asha" not in blob  # body, header, hyperlink targets and metadata
    doc = Document(io.BytesIO(output))
    assert "[EMAIL]" in "\n".join(p.text for p in doc.paragraphs)


def test_docx_redaction_spans_runs_and_keeps_formatting():
    data = SAMPLES["docx"]
    output, _ = redact_document(data, "docx", [RedactItem("+91 90000 00000", "phone")], "removed", False)
    doc = Document(io.BytesIO(output))
    para = next(p for p in doc.paragraphs if "Phone:" in p.text)
    assert para.text.endswith("Phone: [REMOVED]")


def test_unselected_items_are_kept():
    text = "Asha Reddy studied at Example College of Engineering in 2022."
    items = [RedactItem("Asha Reddy", "name")]
    assert redact_text(text, items) == "[REMOVED] studied at Example College of Engineering in 2022."


def test_short_terms_match_whole_words_only():
    assert redact_text("She uses Shell and her other tools", [RedactItem("she", "gender"), RedactItem("her", "gender")]) == (
        "[REMOVED] uses Shell and [REMOVED] other tools"
    )


# --------------------------------------------------------------------------- API

client = TestClient(app)


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_api_round_trip(kind):
    upload = {"file": (f"resume.{kind}", SAMPLES[kind])}
    analysis = client.post("/api/analyze", files=upload)
    assert analysis.status_code == 200
    body = analysis.json()
    assert body["kind"] == kind and body["findings"]

    items = [{"text": f["text"], "category": f["category"]} for f in body["findings"] if f["default_selected"]]
    res = client.post("/api/redact", files=upload, data={"options": json.dumps({"items": items})})
    assert res.status_code == 200
    out = res.json()
    assert out["filename"] == f"resume_anonymized.{kind}"
    assert out["residual"] == []
    assert "asha" not in out["output_text"].lower()
    assert base64.b64decode(out["file_base64"])[:2] in (b"%P", b"PK")


def test_api_rejects_unsupported_files():
    res = client.post("/api/analyze", files={"file": ("notes.txt", b"hello")})
    assert res.status_code == 415
