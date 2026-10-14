"""Real-world resume layouts: placement-cell templates, tables, sidebars, bullets and dates."""
import pytest

from app.detect import detect_pii
from app.extract import extract
from app.profile import build_profile

from .layout_resumes import make_placement_pdf, make_two_column_pdf

PLACEMENT = make_placement_pdf()
TWO_COLUMN = make_two_column_pdf()


def lines_of(data):
    return extract(data, "pdf").text.splitlines()


def profile_of(data):
    ex = extract(data, "pdf")
    return build_profile(ex.text, detect_pii(ex.text), image_count=ex.image_count)


def section(profile, key):
    return next(s for s in profile["sections"] if s["key"] == key)


# --------------------------------------------------------------------------- reading order


def test_name_block_drawn_last_is_read_first():
    assert lines_of(PLACEMENT)[0] == "KAVYA MENON"


def test_letterhead_repeated_on_every_page_is_dropped():
    text = extract(PLACEMENT, "pdf").text
    assert "Training and Placement" not in text and "placements@example.edu" not in text


def test_table_rows_stay_together():
    assert "2021 | Class XII | Lotus Junior College, Madurai, CBSE | 95.20%" in lines_of(PLACEMENT)


def test_bullets_dates_and_wrapped_lines_are_rebuilt():
    lines = lines_of(PLACEMENT)
    assert "• Manufacturing Analytics Internship | Orbit Motors Ltd. (May 2025)" in lines
    # the indented "o" bullets are sub-points of the internship title above them
    assert any(l.startswith("◦ Built a Python model") and l.endswith("within the first month.") for l in lines)
    assert "o" not in lines and "•" not in lines  # no orphaned bullet glyphs


def test_skill_label_and_value_columns_are_joined():
    assert "• Programming Languages: Python, C++" in lines_of(PLACEMENT)


def test_two_column_sidebar_is_read_separately_from_main_column():
    lines = lines_of(TWO_COLUMN)
    assert lines.index("SKILLS") < lines.index("SQL") < lines.index("Python") < lines.index("LANGUAGES")
    assert not any("|" in l for l in lines)  # sidebar text never glued to main-column text
    assert any(l.startswith("Data analyst with two years") and l.endswith("every week.") for l in lines)


# --------------------------------------------------------------------------- blind profile


def test_placement_template_profile_keeps_every_professional_section():
    profile, identity = profile_of(PLACEMENT)
    keys = [s["key"] for s in profile["sections"]]
    for key in ("skills", "experience", "projects", "leadership", "education", "certifications",
                "achievements", "activities"):
        assert key in keys, f"{key} section missing"
    assert profile["skills"] == ["Python", "C++", "SolidWorks", "ANSYS"]
    assert section(profile, "certifications")["lines"] == ["Six Sigma Green Belt"]
    assert "• Head, Robotics Club (Aug 2024 - Present)" in section(profile, "leadership")["lines"]
    assert profile["headline"] == "B.Tech - Mechanical Engineering"
    assert identity["name"] == "Kavya Menon"


def test_placement_template_profile_hides_identity_and_bias_clues():
    profile, _ = profile_of(PLACEMENT)
    shown = "\n".join([profile["headline"]] + [l for s in profile["sections"] for l in s["lines"]]).lower()
    for value in ("kavya", "menon", "female", "14/03/2004", "98765", "example.com",
                  "lotus junior college", "green valley", "example institute", "madurai", "2019", "2021"):
        assert value not in shown, f"{value!r} leaked"
    assert "orbit motors" in shown and "9.10" in shown  # employer and grades stay


def test_two_column_profile():
    profile, identity = profile_of(TWO_COLUMN)
    assert profile["skills"] == ["SQL", "Python", "Tableau", "Excel"]
    assert set(profile["hidden"]["sections"]) == {"Contact", "Languages"}
    experience = section(profile, "experience")["lines"]
    assert len(experience) == 3  # two achievements stay two separate lines
    assert profile["headline"] == "Data Analyst"
    assert identity["phones"] == ["+91 91234 56789"]


@pytest.mark.parametrize("line", ["Sri Krishna Pharmaceuticals Ltd.", "Shri Ram Finance"])
def test_sri_in_company_names_is_not_a_gendered_title(line):
    assert not [f for f in detect_pii(f"Intern at {line}") if f.category == "gender"]


def test_one_word_event_names_are_not_hidden():
    text = "ASHA RAO\nACTIVITIES\n• Won first place in Choreonite at the cultural fest organised by Apekshaa."
    profile, _ = build_profile(text, detect_pii(text))
    assert "Choreonite" in section(profile, "activities")["lines"][0]


def test_docx_tabs_and_list_bullets():
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Data Intern, Acme Corp\tJun 2024 - Aug 2024", style="List Bullet")
    doc.add_paragraph("Cleaned 2M rows of sales data with pandas.", style="List Bullet 2")
    buf = io.BytesIO()
    doc.save(buf)
    lines = extract(buf.getvalue(), "docx").text.splitlines()
    assert "• Data Intern, Acme Corp (Jun 2024 - Aug 2024)" in lines
    assert "◦ Cleaned 2M rows of sales data with pandas." in lines


def test_blind_pdf_renders_tables_and_entries():
    import pymupdf

    from app.blind_pdf import render

    profile, _ = profile_of(PLACEMENT)
    text = pymupdf.open(stream=render("CAND-TEST", profile), filetype="pdf")[0].get_text()
    assert "Orbit Motors" in text and "B.Tech - Mechanical" in text and "Kavya" not in text


def test_action_verbs_at_the_start_of_a_bullet_are_not_names():
    text = ("ASHA RAO\nLEADERSHIP\n• Spearheaded Bag of Delight, organizing collection drives to mobilize "
            "stationery donations and distributing educational supplies to government school children.")
    assert not [f for f in detect_pii(text) if f.category == "name" and "Spearheaded" in f.text]
